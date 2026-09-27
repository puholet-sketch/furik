# -*- coding: utf-8 -*-
"""Rebuild raschety.html as PMO executive report visual."""
from pathlib import Path
from html import escape
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
XLSX = ROOT / "заказы-купер" / "реестр-перевыставление-Фурик.xlsx"


def rub(x: float) -> str:
    return f"{x:,.2f}".replace(",", "\u00a0").replace(".", ",")


def fmt_date(d) -> str:
    y, m, day = str(d)[:10].split("-")
    return f"{day}.{m}.{y}"


wb = load_workbook(XLSX)
ws = wb.active
trs = []
total = 0.0
n = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    d, no, pos, qty, unit, sm, src = r[:7]
    sm = float(sm)
    total += sm
    n += 1
    store = (src or "").replace(" / Купер", "").strip()
    trs.append(
        "<tr>"
        f'<td data-label="Дата">{escape(fmt_date(d))}</td>'
        f'<td data-label="№ заказа"><code>{escape(str(no or ""))}</code></td>'
        f'<td data-label="Товар">{escape(pos or "")}</td>'
        f'<td data-label="Сумма" class="num">{rub(sm)}&nbsp;₽</td>'
        f'<td data-label="Магазин">{escape(store)}</td>'
        "</tr>"
    )
assert n == 40 and abs(total - 9853.56) < 0.02
rows = "\n".join(trs)

html = f"""<!DOCTYPE html>
<html lang="ru">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Сводка взаиморасчётов — Ходжиматов Фуркатжон</title>
  <link rel="stylesheet" href="styles.css" />
  <style>
    /* —— PMO executive report (page-scoped) —— */
    body.page-pmo {{
      --ink: #1a1a1a;
      --text: #333333;
      --muted: #666666;
      --line: #e2e2e2;
      --paper: #f7f6f4;
      --white: #ffffff;
      --accent: #1a1a1a;
      background: var(--paper);
      color: var(--text);
      font-size: 15px;
      line-height: 1.5;
    }}
    body.page-pmo a {{ color: #1a1a1a; text-decoration-color: #999; }}
    body.page-pmo a:hover {{ color: #000; text-decoration-color: #000; }}

    body.page-pmo .topbar {{
      background: rgba(247, 246, 244, 0.96);
      border-bottom: 1px solid var(--line);
      backdrop-filter: none;
    }}
    body.page-pmo .brand__mark {{
      width: 10px; height: 10px; flex: 0 0 10px;
      border-radius: 0;
      background: var(--ink);
      box-shadow: none;
    }}
    body.page-pmo .brand__mark::after {{ display: none; }}
    body.page-pmo .brand__name {{
      font-size: 12px; font-weight: 700; letter-spacing: 0.06em;
      text-transform: uppercase; color: var(--ink);
    }}
    body.page-pmo .brand__sub {{ font-size: 12px; color: var(--muted); }}
    body.page-pmo .topnav a {{
      color: var(--muted); font-size: 13px; text-decoration: none;
      padding: 6px 8px;
    }}
    body.page-pmo .topnav a:hover {{ color: var(--ink); background: transparent; }}
    body.page-pmo .nav-toggle {{
      border-radius: 2px; border-color: var(--line); background: var(--white);
    }}

    body.page-pmo main {{
      max-width: 920px;
      margin: 0 auto;
      padding: 28px 20px 48px;
    }}

    body.page-pmo .report-head {{
      border-bottom: 1px solid var(--ink);
      padding-bottom: 20px;
      margin-bottom: 20px;
    }}
    body.page-pmo .report-head .eyebrow {{
      margin: 0 0 8px;
      font-size: 11px;
      letter-spacing: 0.08em;
      text-transform: uppercase;
      color: var(--muted);
      font-weight: 600;
    }}
    body.page-pmo .report-head h1 {{
      margin: 0 0 10px;
      font-size: clamp(1.35rem, 2.5vw, 1.75rem);
      font-weight: 700;
      color: var(--ink);
      letter-spacing: -0.01em;
      line-height: 1.25;
      max-width: none;
    }}
    body.page-pmo .report-head .lead {{
      margin: 0;
      max-width: 44rem;
      color: var(--muted);
      font-size: 0.95rem;
    }}

    body.page-pmo .meta-strip {{
      display: grid;
      grid-template-columns: repeat(4, minmax(0, 1fr));
      gap: 0;
      border: 1px solid var(--line);
      background: var(--white);
      margin-bottom: 28px;
    }}
    body.page-pmo .meta-strip__item {{
      padding: 14px 16px;
      border-right: 1px solid var(--line);
    }}
    body.page-pmo .meta-strip__item:last-child {{ border-right: 0; }}
    body.page-pmo .meta-strip__label {{
      display: block;
      font-size: 11px;
      letter-spacing: 0.04em;
      text-transform: uppercase;
      color: var(--muted);
      margin-bottom: 6px;
    }}
    body.page-pmo .meta-strip__value {{
      display: block;
      font-size: 1.15rem;
      font-weight: 700;
      color: var(--ink);
      font-variant-numeric: tabular-nums;
      letter-spacing: -0.01em;
    }}
    body.page-pmo .meta-strip__hint {{
      display: block;
      margin-top: 4px;
      font-size: 12px;
      color: var(--muted);
    }}

    body.page-pmo .report-section {{
      background: transparent;
      border: 0;
      box-shadow: none;
      border-radius: 0;
      padding: 0 0 28px;
      margin: 0 0 8px;
      border-bottom: 1px solid var(--line);
    }}
    body.page-pmo .report-section:last-of-type {{ border-bottom: 0; }}
    body.page-pmo .report-section .sec-head {{
      display: flex;
      gap: 14px;
      align-items: baseline;
      margin: 0 0 10px;
      padding-top: 8px;
    }}
    body.page-pmo .report-section .sec-num {{
      flex: 0 0 auto;
      font-size: 13px;
      font-weight: 700;
      color: var(--ink);
      letter-spacing: 0.04em;
      min-width: 1.6em;
    }}
    body.page-pmo .report-section h2 {{
      margin: 0;
      font-size: 1.05rem;
      font-weight: 700;
      color: var(--ink);
      letter-spacing: -0.01em;
    }}
    body.page-pmo .report-section .sec-sub {{
      margin: 2px 0 0;
      font-size: 12px;
      color: var(--muted);
    }}
    body.page-pmo .section-intro {{
      margin: 0 0 14px;
      color: var(--muted);
      font-size: 0.92rem;
      max-width: 46rem;
    }}

    body.page-pmo .table-wrap {{
      border: 1px solid var(--line);
      border-radius: 0;
      background: var(--white);
      overflow-x: auto;
      -webkit-overflow-scrolling: touch;
    }}
    body.page-pmo table {{
      width: 100%;
      border-collapse: collapse;
      font-size: 13px;
      table-layout: fixed;
    }}
    body.page-pmo thead th {{
      background: #f0efed;
      color: var(--ink);
      font-weight: 600;
      font-size: 11px;
      letter-spacing: 0.03em;
      text-transform: uppercase;
      text-align: left;
      padding: 10px 12px;
      border-bottom: 1px solid var(--line);
    }}
    body.page-pmo tbody td {{
      padding: 9px 12px;
      border-bottom: 1px solid var(--line);
      vertical-align: top;
      color: var(--text);
    }}
    body.page-pmo tbody tr:last-child td {{ border-bottom: 0; }}
    body.page-pmo .num,
    body.page-pmo th.num {{
      text-align: right !important;
      font-variant-numeric: tabular-nums;
      white-space: nowrap;
    }}
    body.page-pmo .row-total td {{
      background: #f0efed;
      font-weight: 700;
      color: var(--ink);
    }}
    body.page-pmo code {{
      font-size: 0.85em;
      background: #f0efed;
      color: var(--ink);
      padding: 0.1em 0.35em;
      border-radius: 2px;
      white-space: nowrap;
    }}

    body.page-pmo .badge {{
      display: inline-block;
      font-size: 11px;
      font-weight: 600;
      letter-spacing: 0.02em;
      padding: 3px 8px;
      border-radius: 2px;
      border: 1px solid transparent;
      white-space: nowrap;
    }}
    body.page-pmo .badge--ok {{
      color: #1f5c38;
      background: #eef6f1;
      border-color: #c9dfd2;
    }}
    body.page-pmo .badge--pending {{
      color: #5c5c5c;
      background: #f3f3f3;
      border-color: #dedede;
    }}

    body.page-pmo .note-box {{
      margin-top: 12px;
      padding: 12px 14px;
      border-left: 2px solid var(--ink);
      background: var(--white);
      border-top: 1px solid var(--line);
      border-right: 1px solid var(--line);
      border-bottom: 1px solid var(--line);
      font-size: 0.9rem;
      color: var(--muted);
    }}
    body.page-pmo .note {{
      margin-top: 12px;
      font-size: 0.88rem;
      color: var(--muted);
    }}

    body.page-pmo .doc-actions {{
      display: flex;
      flex-wrap: wrap;
      gap: 8px;
    }}
    body.page-pmo .btn {{
      border-radius: 2px;
      font-weight: 600;
      font-size: 13px;
      padding: 9px 14px;
      box-shadow: none;
      text-decoration: none;
      border: 1px solid var(--line);
      background: var(--white);
      color: var(--ink);
      transition: none;
    }}
    body.page-pmo .btn:hover {{
      border-color: var(--ink);
      background: var(--white);
      color: var(--ink);
      transform: none;
    }}
    body.page-pmo .btn--primary,
    body.page-pmo .btn--pdf {{
      background: var(--ink);
      border-color: var(--ink);
      color: #fff;
    }}
    body.page-pmo .btn--primary:hover,
    body.page-pmo .btn--pdf:hover {{
      background: #000;
      border-color: #000;
      color: #fff;
    }}
    body.page-pmo .btn--ghost {{
      background: var(--white);
      color: var(--ink);
    }}

    body.page-pmo .table-registry {{ min-width: 700px; }}
    body.page-pmo .table-registry col.c-date {{ width: 11%; }}
    body.page-pmo .table-registry col.c-order {{ width: 16%; }}
    body.page-pmo .table-registry col.c-item {{ width: 43%; }}
    body.page-pmo .table-registry col.c-sum {{ width: 14%; }}
    body.page-pmo .table-registry col.c-store {{ width: 16%; }}

    body.page-pmo .page-foot {{
      margin-top: 24px;
      padding-top: 16px;
      border-top: 1px solid var(--line);
      font-size: 12px;
      color: var(--muted);
    }}

    @media (max-width: 760px) {{
      body.page-pmo .meta-strip {{ grid-template-columns: 1fr 1fr; }}
      body.page-pmo .meta-strip__item:nth-child(2) {{ border-right: 0; }}
      body.page-pmo .meta-strip__item:nth-child(1),
      body.page-pmo .meta-strip__item:nth-child(2) {{
        border-bottom: 1px solid var(--line);
      }}
      body.page-pmo .btn {{ width: auto; }}
    }}
    @media (max-width: 480px) {{
      body.page-pmo .meta-strip {{ grid-template-columns: 1fr; }}
      body.page-pmo .meta-strip__item {{ border-right: 0; border-bottom: 1px solid var(--line); }}
      body.page-pmo .meta-strip__item:last-child {{ border-bottom: 0; }}
    }}
    @media (max-width: 640px) {{
      body.page-pmo .table-stack thead {{ display: none; }}
      body.page-pmo .table-stack tr {{
        display: block;
        border-bottom: 1px solid var(--line);
        padding: 10px 0;
      }}
      body.page-pmo .table-stack td {{
        display: flex;
        justify-content: space-between;
        gap: 12px;
        border: 0 !important;
        padding: 4px 12px;
        text-align: right !important;
      }}
      body.page-pmo .table-stack td::before {{
        content: attr(data-label);
        font-weight: 600;
        color: var(--muted);
        text-align: left;
        flex: 0 0 42%;
        text-transform: none;
        letter-spacing: 0;
        font-size: 12px;
      }}
      body.page-pmo .table-registry {{ min-width: 0; }}
    }}
  </style>
</head>
<body class="page-pmo">
  <header class="topbar" id="topbar">
    <div class="topbar__inner">
      <div class="brand">
        <span class="brand__mark" aria-hidden="true"></span>
        <div>
          <div class="brand__name">Сводка ПМО</div>
          <div class="brand__sub">Ходжиматов Фуркатжон · 2026</div>
        </div>
      </div>
      <button class="nav-toggle" type="button" aria-label="Меню" aria-expanded="false" id="navToggle">
        <span></span>
      </button>
      <nav class="topnav" aria-label="Разделы" id="topnav">
        <a href="index.html">Инструкция</a>
        <a href="#zp">1</a>
        <a href="#paid">2</a>
        <a href="#contract">3</a>
        <a href="#registry">4</a>
      </nav>
    </div>
  </header>

  <main>
    <header class="report-head">
      <p class="eyebrow">ИП Сорванова · служебная сводка · 27.09.2026</p>
      <h1>Расчёты по сотруднику Ходжиматову Фуркатжону</h1>
      <p class="lead">
        Официальная оплата труда (0,125 ставки с 25.08.2026), платежи в бюджет по выписке
        и реестр личных покупок к перевыставлению.
      </p>
    </header>

    <div class="meta-strip" aria-label="Ключевые показатели">
      <div class="meta-strip__item">
        <span class="meta-strip__label">На руки / мес.</span>
        <span class="meta-strip__value">4&nbsp;320,25&nbsp;₽</span>
        <span class="meta-strip__hint">оклад 4&nbsp;966,25 − НДФЛ</span>
      </div>
      <div class="meta-strip__item">
        <span class="meta-strip__label">Расход ИП / мес.</span>
        <span class="meta-strip__value">6&nbsp;466,06&nbsp;₽</span>
        <span class="meta-strip__hint">оклад + взносы + травматизм</span>
      </div>
      <div class="meta-strip__item">
        <span class="meta-strip__label">Налоги и взносы</span>
        <span class="meta-strip__value">2&nbsp;145,81&nbsp;₽</span>
        <span class="meta-strip__hint">НДФЛ + 30% + 0,2%</span>
      </div>
      <div class="meta-strip__item">
        <span class="meta-strip__label">Реестр покупок</span>
        <span class="meta-strip__value">9&nbsp;853,56&nbsp;₽</span>
        <span class="meta-strip__hint">40 позиций · 2026</span>
      </div>
    </div>

    <section id="zp" class="report-section">
      <div class="sec-head">
        <span class="sec-num">1.</span>
        <div>
          <h2>Официальная оплата труда (0,125 ставки)</h2>
          <p class="sec-sub">Ходжиматов Фуркатжон Махамаджонович · допсоглашение с 25.08.2026</p>
        </div>
      </div>
      <p class="section-intro">
        Официальная часть — 1/8 полной ставки (оклад МЗП Москвы 39&nbsp;730&nbsp;₽).
        Структура начисления, удержаний и обязательных платежей работодателя за полный месяц.
      </p>

      <div class="table-wrap">
        <table>
          <thead>
            <tr>
              <th>Параметр</th>
              <th class="num">Полная ставка</th>
              <th class="num">0,125 ставки</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td data-label="Параметр">Оклад (начисление)</td>
              <td data-label="Полная ставка" class="num">39&nbsp;730,00&nbsp;₽</td>
              <td data-label="0,125 ставки" class="num"><strong>4&nbsp;966,25&nbsp;₽</strong></td>
            </tr>
            <tr>
              <td data-label="Параметр">Норма рабочего времени</td>
              <td data-label="Полная ставка" class="num">40 ч / нед.</td>
              <td data-label="0,125 ставки" class="num">5 ч / нед. (1 ч / день)</td>
            </tr>
          </tbody>
        </table>
      </div>

      <div class="table-wrap" style="margin-top:12px">
        <table>
          <thead>
            <tr>
              <th>Статья</th>
              <th class="num">Сумма</th>
              <th>Куда / кто</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td data-label="Статья">Выплата сотруднику</td>
              <td data-label="Сумма" class="num"><strong>4&nbsp;320,25&nbsp;₽</strong></td>
              <td data-label="Куда">на счёт в ПАО Сбербанк</td>
            </tr>
            <tr>
              <td data-label="Статья">НДФЛ 13%</td>
              <td data-label="Сумма" class="num">646,00&nbsp;₽</td>
              <td data-label="Куда">удержание из оклада → ФНС (ЕНП)</td>
            </tr>
            <tr>
              <td data-label="Статья">Страховые взносы 30%</td>
              <td data-label="Сумма" class="num">1&nbsp;489,88&nbsp;₽</td>
              <td data-label="Куда">за счёт ИП → ФНС (ЕНП)</td>
            </tr>
            <tr>
              <td data-label="Статья">Взносы на травматизм 0,2%</td>
              <td data-label="Сумма" class="num">9,93&nbsp;₽</td>
              <td data-label="Куда">за счёт ИП → СФР</td>
            </tr>
            <tr class="row-total">
              <td data-label="Статья">Итого НДФЛ + взносы + травматизм</td>
              <td data-label="Сумма" class="num"><strong>2&nbsp;145,81&nbsp;₽</strong></td>
              <td data-label="Куда">в месяц при 0,125 ставки</td>
            </tr>
          </tbody>
        </table>
      </div>

      <div class="note-box">
        Совокупный расход ИП по официальной части:
        4&nbsp;966,25&nbsp;₽ (оклад) + 1&nbsp;489,88&nbsp;₽ + 9,93&nbsp;₽ =
        <strong>6&nbsp;466,06&nbsp;₽</strong> в месяц. НДФЛ входит в оклад и сверху не добавляется.
        Суммы сверх официальной части учитываются отдельно в рамках утверждённого недельного бюджета проекта
        и в настоящую таблицу не включены.
      </div>
    </section>

    <section id="paid" class="report-section">
      <div class="sec-head">
        <span class="sec-num">2.</span>
        <div>
          <h2>Уплачено в бюджет (по выписке)</h2>
          <p class="sec-sub">Подтверждённые платежи по Фуркатжону · август–сентябрь 2026</p>
        </div>
      </div>
      <p class="section-intro">
        Строки сверены с банковской выпиской. Указаны только фактически проведённые платежи;
        плановые суммы с ненаступившим сроком — в примечании.
      </p>

      <div class="table-wrap">
        <table>
          <thead>
            <tr>
              <th>Дата</th>
              <th class="num">Сумма</th>
              <th>Назначение</th>
              <th>Период</th>
              <th>Зачтено</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td data-label="Дата">23.08.2026</td>
              <td data-label="Сумма" class="num">2&nbsp;214,00&nbsp;₽</td>
              <td data-label="Назначение">НДФЛ → ФНС (ЕНП)</td>
              <td data-label="Период">август 2026 (часть)</td>
              <td data-label="Зачтено"><span class="badge badge--ok">уже зачтено</span></td>
            </tr>
            <tr>
              <td data-label="Дата">23.09.2026</td>
              <td data-label="Сумма" class="num">1&nbsp;137,00&nbsp;₽</td>
              <td data-label="Назначение">НДФЛ → ФНС (ЕНП)</td>
              <td data-label="Период">август 2026 (доплата)</td>
              <td data-label="Зачтено"><span class="badge badge--pending">еще не зачтено</span></td>
            </tr>
            <tr>
              <td data-label="Дата">23.09.2026</td>
              <td data-label="Сумма" class="num">323,00&nbsp;₽</td>
              <td data-label="Назначение">НДФЛ → ФНС (ЕНП)</td>
              <td data-label="Период">сентябрь 2026 (аванс)</td>
              <td data-label="Зачтено"><span class="badge badge--pending">еще не зачтено</span></td>
            </tr>
            <tr>
              <td data-label="Дата">23.09.2026</td>
              <td data-label="Сумма" class="num">7&nbsp;733,16&nbsp;₽</td>
              <td data-label="Назначение">Взносы 30% → ФНС (ЕНП)</td>
              <td data-label="Период">август 2026</td>
              <td data-label="Зачтено"><span class="badge badge--pending">еще не зачтено</span></td>
            </tr>
          </tbody>
        </table>
      </div>

      <p class="note">
        Август: НДФЛ итого <strong>3&nbsp;351,00&nbsp;₽</strong> (2&nbsp;214,00 + 1&nbsp;137,00) и взносы
        <strong>7&nbsp;733,16&nbsp;₽</strong> — соответствуют начислению 25&nbsp;777,20&nbsp;₽.
        Сентябрь: НДФЛ с аванса <strong>323,00&nbsp;₽</strong> уплачен; доплата НДФЛ ≈ 323,00&nbsp;₽ и взносы
        1&nbsp;489,88&nbsp;₽ — срок до 28.10.2026.
        Платежи на травматизм 0,2% в СФР (август ≈ 51,55&nbsp;₽, сентябрь 9,93&nbsp;₽) в данной выписке
        не подтверждены и требуют отдельной проверки.
      </p>
    </section>

    <section id="contract" class="report-section">
      <div class="sec-head">
        <span class="sec-num">3.</span>
        <div>
          <h2>Документы</h2>
          <p class="sec-sub">Трудовой договор от 06.08.2026 и допсоглашение об 0,125 ставки</p>
        </div>
      </div>
      <p class="section-intro">
        Канонические файлы для кадрового учёта и бухгалтерии. При расхождении версий приоритет у PDF
        с подписями сторон.
      </p>
      <div class="doc-actions">
        <a class="btn btn--pdf" href="трудовой-договор/трудовой-договор-Ходжиматов-бессрочный-06.08.2026.pdf" download>Трудовой договор (PDF)</a>
        <a class="btn btn--ghost" href="трудовой-договор/трудовой-договор-Ходжиматов-бессрочный-06.08.2026.docx" download>Трудовой договор (Word)</a>
        <a class="btn btn--ghost" href="трудовой-договор/допсоглашение-0.125-ставки-Ходжиматов.pdf" download>Допсоглашение 0,125 (PDF)</a>
      </div>
    </section>

    <section id="registry" class="report-section">
      <div class="sec-head">
        <span class="sec-num">4.</span>
        <div>
          <h2>Реестр личных покупок из Сбермаркета</h2>
          <p class="sec-sub">К перевыставлению Фуркатжону · итого <strong>9&nbsp;853,56&nbsp;₽</strong> · 40 позиций</p>
        </div>
      </div>
      <p class="section-intro">
        Согласованные личные позиции из заказов доставки Сбермаркет.
        Товары кофейни, цитрусы и сервисные сборы не учитываются.
      </p>
      <div class="doc-actions" style="margin-bottom:12px">
        <a class="btn btn--ghost" href="заказы-купер/реестр-перевыставление-Фурик.xlsx" download>Скачать реестр (Excel)</a>
      </div>
      <div class="table-wrap">
        <table class="table-stack table-registry">
          <colgroup>
            <col class="c-date" /><col class="c-order" /><col class="c-item" /><col class="c-sum" /><col class="c-store" />
          </colgroup>
          <thead>
            <tr>
              <th>Дата</th>
              <th>№ заказа</th>
              <th>Товар</th>
              <th class="num">Сумма</th>
              <th>Магазин</th>
            </tr>
          </thead>
          <tbody>
{rows}
            <tr class="row-total">
              <td data-label="Итого" colspan="3"><strong>Итого к перевыставлению</strong></td>
              <td data-label="Сумма" class="num"><strong>9&nbsp;853,56&nbsp;₽</strong></td>
              <td data-label="Позиции">40 позиций</td>
            </tr>
          </tbody>
        </table>
      </div>
    </section>

    <p class="page-foot">
      Сопутствующий материал: <a href="index.html">инструкция по приёму сотрудника</a>.
      Суммы по оплате труда — расчётный ориентир; перед платежами рекомендуется сверка с бухгалтером.
    </p>
  </main>

  <script>
    (function () {{
      const toggle = document.getElementById("navToggle");
      const nav = document.getElementById("topnav");
      if (!toggle || !nav) return;
      toggle.addEventListener("click", () => {{
        const open = nav.classList.toggle("is-open");
        toggle.setAttribute("aria-expanded", open ? "true" : "false");
      }});
      nav.querySelectorAll("a").forEach((a) => {{
        a.addEventListener("click", () => {{
          nav.classList.remove("is-open");
          toggle.setAttribute("aria-expanded", "false");
        }});
      }});
    }})();
  </script>
</body>
</html>
"""

(DOCS / "raschety.html").write_text(html, encoding="utf-8")
# Keep rebuild script in sync as the canonical builder
(ROOT / "_tools" / "rebuild_raschety_pro.py").write_text(
    Path(__file__).read_text(encoding="utf-8"), encoding="utf-8"
)
print(f"ok rows={n} total={total:.2f} bytes={len(html)}")
