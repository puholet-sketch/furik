# -*- coding: utf-8 -*-
"""Rebuild raschety.html in project-office (black/gold) visual system."""
from pathlib import Path
from html import escape
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
XLSX = ROOT / "заказы-купер" / "реестр-перевыставление-Фурик.xlsx"


def rub(x: float) -> str:
    return f"{x:,.2f}".replace(",", "\u00a0").replace(".", ",")


def fmt_date(d) -> str:
    y, m, day = str(d)[:10].split("-")
    return f"{day}.{m}.{y}"


wb = load_workbook(XLSX)
ws = wb.active
trs = []
total = 0.0
n = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    d, no, pos, qty, unit, sm, src = r[:7]
    sm = float(sm)
    total += sm
    n += 1
    store = (src or "").replace(" / Купер", "").strip()
    trs.append(
        "<tr>"
        f'<td data-label="Дата">{escape(fmt_date(d))}</td>'
        f'<td data-label="№ заказа"><code>{escape(str(no or ""))}</code></td>'
        f'<td data-label="Товар">{escape(pos or "")}</td>'
        f'<td data-label="Сумма" class="num">{rub(sm)}&nbsp;₽</td>'
        f'<td data-label="Магазин">{escape(store)}</td>'
        "</tr>"
    )
assert n == 40 and abs(total - 9853.56) < 0.02
rows = "\n".join(trs)

html = f"""<!DOCTYPE html>
<html lang="ru">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Сводка взаиморасчётов — Ходжиматов Фуркатжон</title>
  <link rel="preconnect" href="https://fonts.googleapis.com" />
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin />
  <link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Outfit:wght@600;700&display=swap" rel="stylesheet" />
  <style>
    :root {{
      --po-bg: #000000;
      --po-panel: #111111;
      --po-panel-alt: #161616;
      --po-gold: #f0c419;
      --po-gold-deep: #d4a80f;
      --po-ink: #ffffff;
      --po-muted: #8a8a8a;
      --po-body: #b8b8b8;
      --po-line: #2a2a2a;
      --po-line-soft: #1a1a1a;
      --po-radius: 0.85rem;
      --font-display: "Outfit", "Segoe UI", sans-serif;
      --font-body: "Inter", "Segoe UI", sans-serif;
      --max: 1100px;
    }}
    *, *::before, *::after {{ box-sizing: border-box; }}
    html {{ scroll-behavior: smooth; -webkit-text-size-adjust: 100%; }}
    body {{
      margin: 0;
      min-height: 100%;
      background: var(--po-bg);
      color: var(--po-ink);
      font-family: var(--font-body);
      font-size: 15px;
      line-height: 1.55;
    }}
    a {{ color: var(--po-ink); text-decoration-color: #555; }}
    a:hover {{ color: var(--po-gold); text-decoration-color: var(--po-gold); }}

    .topbar {{
      position: sticky; top: 0; z-index: 40;
      border-bottom: 1px solid #1f1f1f;
      background: rgba(0, 0, 0, 0.96);
      backdrop-filter: blur(12px);
    }}
    .topbar__inner {{
      max-width: var(--max);
      margin: 0 auto;
      padding: 10px 16px;
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 12px;
    }}
    .brand {{ display: flex; flex-direction: column; gap: 2px; min-width: 0; }}
    .brand__name {{
      font-family: var(--font-display);
      font-size: 0.85rem;
      font-weight: 700;
      color: var(--po-ink);
    }}
    .brand__name em {{ font-style: normal; color: var(--po-gold); }}
    .brand__sub {{ font-size: 0.72rem; color: var(--po-muted); }}
    .topnav {{ display: flex; flex-wrap: wrap; gap: 2px; }}
    .topnav a {{
      border-radius: 6px;
      padding: 6px 10px;
      font-size: 0.75rem;
      font-weight: 500;
      color: var(--po-muted);
      text-decoration: none;
    }}
    .topnav a:hover {{ background: #1c1c1c; color: var(--po-ink); }}
    .nav-toggle {{
      display: none;
      width: 40px; height: 40px;
      border: 1px solid var(--po-line);
      border-radius: 6px;
      background: #111;
      cursor: pointer;
      align-items: center; justify-content: center;
    }}
    .nav-toggle span,
    .nav-toggle span::before,
    .nav-toggle span::after {{
      display: block; width: 16px; height: 1.5px; background: var(--po-ink); position: relative;
    }}
    .nav-toggle span::before,
    .nav-toggle span::after {{ content: ""; position: absolute; left: 0; }}
    .nav-toggle span::before {{ top: -5px; }}
    .nav-toggle span::after {{ top: 5px; }}

    .sheet {{
      max-width: var(--max);
      margin: 0 auto;
      padding: 1.75rem 16px 4rem;
    }}
    .po-back {{
      display: inline-block;
      font-size: 0.875rem;
      color: var(--po-muted);
      text-decoration: none;
      margin-bottom: 1.25rem;
    }}
    .po-back:hover {{ color: var(--po-gold); }}

    .po-hero {{
      position: relative;
      overflow: hidden;
      padding: 1.75rem 0 1.5rem;
      border-bottom: 1px solid var(--po-line-soft);
      margin-bottom: 0.5rem;
    }}
    .po-kicker {{
      margin: 0 0 0.65rem;
      font-family: var(--font-display);
      font-size: 0.72rem;
      font-weight: 600;
      letter-spacing: 0.16em;
      text-transform: uppercase;
      color: var(--po-gold);
    }}
    .po-hero__title {{
      margin: 0;
      width: 100%;
      max-width: none;
      font-family: var(--font-display);
      font-size: clamp(1.75rem, 5vw, 2.85rem);
      font-weight: 700;
      line-height: 1.05;
      letter-spacing: 0.02em;
      text-transform: uppercase;
      color: #fff;
    }}

    .po-block {{
      margin-top: 2.35rem;
    }}
    .po-block__head {{
      display: flex;
      gap: 0.85rem;
      align-items: flex-start;
      margin-bottom: 1rem;
    }}
    .po-num {{
      flex-shrink: 0;
      display: inline-flex;
      align-items: center;
      justify-content: center;
      min-width: 2rem;
      height: 2rem;
      padding: 0 0.45rem;
      border-radius: 2px;
      background: var(--po-gold);
      font-family: var(--font-display);
      font-size: 0.68rem;
      font-weight: 700;
      letter-spacing: 0.04em;
      color: #000;
    }}
    .po-h2 {{
      margin: 0;
      font-family: var(--font-display);
      font-size: clamp(1.05rem, 2.2vw, 1.35rem);
      line-height: 1.2;
      font-weight: 700;
      letter-spacing: 0.04em;
      text-transform: uppercase;
      color: #fff;
    }}
    .po-sub {{
      margin: 0.35rem 0 0;
      font-size: 0.82rem;
      color: var(--po-muted);
    }}
    .po-lead {{
      margin: 0 0 1rem;
      max-width: 42rem;
      font-size: 0.9rem;
      line-height: 1.55;
      color: var(--po-muted);
    }}

    .table-wrap {{
      border: 1px solid var(--po-line);
      border-radius: var(--po-radius);
      background: var(--po-panel);
      overflow-x: auto;
      -webkit-overflow-scrolling: touch;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      font-size: 0.84rem;
    }}
    thead th {{
      text-align: left;
      padding: 0.7rem 0.85rem;
      background: var(--po-panel-alt);
      color: var(--po-muted);
      font-size: 0.68rem;
      font-weight: 600;
      letter-spacing: 0.08em;
      text-transform: uppercase;
      border-bottom: 1px solid var(--po-line);
    }}
    tbody td {{
      padding: 0.65rem 0.85rem;
      border-bottom: 1px solid var(--po-line-soft);
      color: var(--po-body);
      vertical-align: top;
    }}
    tbody tr:last-child td {{ border-bottom: 0; }}
    .num, th.num {{
      text-align: right !important;
      font-variant-numeric: tabular-nums;
      white-space: nowrap;
      color: var(--po-ink);
    }}
    .row-total td {{
      background: #141414;
      color: var(--po-ink);
      font-weight: 700;
    }}
    code {{
      font-family: ui-monospace, Consolas, monospace;
      font-size: 0.8em;
      color: var(--po-gold);
      background: rgba(240, 196, 25, 0.08);
      padding: 0.1em 0.35em;
      border-radius: 3px;
      white-space: nowrap;
    }}

    .badge {{
      display: inline-block;
      font-size: 0.7rem;
      font-weight: 600;
      letter-spacing: 0.02em;
      padding: 0.2rem 0.5rem;
      border-radius: 3px;
      white-space: nowrap;
    }}
    .badge--ok {{
      color: #000;
      background: var(--po-gold);
    }}
    .badge--pending {{
      color: var(--po-muted);
      background: #1c1c1c;
      border: 1px solid var(--po-line);
    }}

    .po-notes {{
      display: grid;
      gap: 0.65rem;
      margin-top: 1rem;
    }}
    .po-note {{
      border: 1px solid var(--po-line);
      border-radius: var(--po-radius);
      background: var(--po-panel);
      padding: 0.85rem 1rem;
    }}
    .po-note__label {{
      margin: 0 0 0.35rem;
      font-family: var(--font-display);
      font-size: 0.72rem;
      font-weight: 700;
      letter-spacing: 0.1em;
      text-transform: uppercase;
      color: var(--po-gold);
    }}
    .po-note__text {{
      margin: 0;
      font-size: 0.88rem;
      line-height: 1.5;
      color: var(--po-body);
    }}

    .doc-actions {{
      display: flex;
      flex-wrap: wrap;
      gap: 0.65rem;
    }}
    .btn {{
      display: inline-flex;
      align-items: center;
      justify-content: center;
      padding: 0.65rem 1rem;
      border-radius: 6px;
      font-size: 0.82rem;
      font-weight: 600;
      text-decoration: none;
      border: 1px solid var(--po-line);
      transition: border-color 0.15s ease, background 0.15s ease, color 0.15s ease;
    }}
    .btn--primary {{
      background: var(--po-gold);
      border-color: var(--po-gold);
      color: #000;
    }}
    .btn--primary:hover {{
      background: var(--po-gold-deep);
      border-color: var(--po-gold-deep);
      color: #000;
    }}
    .btn--ghost {{
      background: transparent;
      color: var(--po-ink);
    }}
    .btn--ghost:hover {{
      border-color: var(--po-gold);
      color: var(--po-gold);
    }}

    .table-registry {{ min-width: 700px; }}
    .page-foot {{
      margin-top: 2.5rem;
      padding-top: 1rem;
      border-top: 1px solid var(--po-line-soft);
      font-size: 0.78rem;
      color: var(--po-muted);
    }}

    @media (max-width: 720px) {{
      .nav-toggle {{ display: inline-flex; }}
      .topnav {{
        display: none;
        position: absolute;
        right: 16px; top: 54px;
        flex-direction: column;
        background: #111;
        border: 1px solid var(--po-line);
        border-radius: 8px;
        padding: 6px;
        min-width: 10rem;
      }}
      .topnav.is-open {{ display: flex; }}
      .table-stack thead {{ display: none; }}
      .table-stack tr {{
        display: block;
        border-bottom: 1px solid var(--po-line);
        padding: 0.55rem 0;
      }}
      .table-stack td {{
        display: flex;
        justify-content: space-between;
        gap: 0.75rem;
        border: 0 !important;
        padding: 0.25rem 0.85rem;
        text-align: right !important;
      }}
      .table-stack td::before {{
        content: attr(data-label);
        color: var(--po-muted);
        font-size: 0.72rem;
        font-weight: 600;
        text-align: left;
        flex: 0 0 40%;
        text-transform: uppercase;
        letter-spacing: 0.04em;
      }}
      .table-registry {{ min-width: 0; }}
    }}
  </style>
</head>
<body>
  <header class="topbar" id="topbar">
    <div class="topbar__inner">
      <div class="brand">
        <div class="brand__name">Сводка <em>ПМО</em></div>
        <div class="brand__sub">Ходжиматов Фуркатжон · 2026</div>
      </div>
      <button class="nav-toggle" type="button" aria-label="Меню" aria-expanded="false" id="navToggle"><span></span></button>
      <nav class="topnav" aria-label="Разделы" id="topnav">
        <a href="index.html">Инструкция</a>
        <a href="#zp">01</a>
        <a href="#paid">02</a>
        <a href="#contract">03</a>
        <a href="#registry">04</a>
      </nav>
    </div>
  </header>

  <main class="sheet">
    <a class="po-back" href="index.html">← Инструкция по приёму</a>

    <header class="po-hero">
      <p class="po-kicker">ИП Сорванова · служебная сводка · 27.09.2026</p>
      <h1 class="po-hero__title">Расчёты по сотруднику Ходжиматову Фуркатжону</h1>
    </header>

    <section id="zp" class="po-block">
      <div class="po-block__head">
        <span class="po-num">01</span>
        <div>
          <h2 class="po-h2">Официальная оплата труда (0,125 ставки)</h2>
          <p class="po-sub">Ходжиматов Фуркатжон Махамаджонович · допсоглашение с 25.08.2026</p>
        </div>
      </div>
      <p class="po-lead">
        Официальная часть — 1/8 полной ставки (оклад МЗП Москвы 39&nbsp;730&nbsp;₽).
        Структура начисления, удержаний и обязательных платежей работодателя за полный месяц.
        Суммы сверх официальной части учитываются отдельно в рамках недельного бюджета проекта.
      </p>

      <div class="table-wrap">
        <table>
          <thead>
            <tr>
              <th>Параметр</th>
              <th class="num">Полная ставка</th>
              <th class="num">0,125 ставки</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td data-label="Параметр">Оклад (начисление)</td>
              <td data-label="Полная ставка" class="num">39&nbsp;730,00&nbsp;₽</td>
              <td data-label="0,125 ставки" class="num"><strong>4&nbsp;966,25&nbsp;₽</strong></td>
            </tr>
            <tr>
              <td data-label="Параметр">Норма рабочего времени</td>
              <td data-label="Полная ставка" class="num">40 ч / нед.</td>
              <td data-label="0,125 ставки" class="num">5 ч / нед. (1 ч / день)</td>
            </tr>
          </tbody>
        </table>
      </div>

      <div class="table-wrap" style="margin-top:0.75rem">
        <table>
          <thead>
            <tr>
              <th>Статья</th>
              <th class="num">Сумма</th>
              <th>Куда / кто</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td data-label="Статья">Выплата сотруднику</td>
              <td data-label="Сумма" class="num"><strong>4&nbsp;320,25&nbsp;₽</strong></td>
              <td data-label="Куда">на счёт в ПАО Сбербанк</td>
            </tr>
            <tr>
              <td data-label="Статья">НДФЛ 13%</td>
              <td data-label="Сумма" class="num">646,00&nbsp;₽</td>
              <td data-label="Куда">удержание из оклада → ФНС (ЕНП)</td>
            </tr>
            <tr>
              <td data-label="Статья">Страховые взносы 30%</td>
              <td data-label="Сумма" class="num">1&nbsp;489,88&nbsp;₽</td>
              <td data-label="Куда">за счёт ИП → ФНС (ЕНП)</td>
            </tr>
            <tr>
              <td data-label="Статья">Взносы на травматизм 0,2%</td>
              <td data-label="Сумма" class="num">9,93&nbsp;₽</td>
              <td data-label="Куда">за счёт ИП → СФР</td>
            </tr>
            <tr class="row-total">
              <td data-label="Статья">Итого НДФЛ + взносы + травматизм</td>
              <td data-label="Сумма" class="num"><strong>2&nbsp;145,81&nbsp;₽</strong></td>
              <td data-label="Куда">в месяц при 0,125 ставки</td>
            </tr>
            <tr class="row-total">
              <td data-label="Статья">Расход ИП (оклад + взносы + травматизм)</td>
              <td data-label="Сумма" class="num"><strong>6&nbsp;466,06&nbsp;₽</strong></td>
              <td data-label="Куда">НДФЛ внутри оклада</td>
            </tr>
          </tbody>
        </table>
      </div>
    </section>

    <section id="paid" class="po-block">
      <div class="po-block__head">
        <span class="po-num">02</span>
        <div>
          <h2 class="po-h2">Уплачено в бюджет (по выписке)</h2>
          <p class="po-sub">Подтверждённые платежи по Фуркатжону · август–сентябрь 2026</p>
        </div>
      </div>
      <p class="po-lead">
        Строки сверены с банковской выпиской. Указаны только фактически проведённые платежи.
      </p>

      <div class="table-wrap">
        <table>
          <thead>
            <tr>
              <th>Дата</th>
              <th class="num">Сумма</th>
              <th>Назначение</th>
              <th>Период</th>
              <th>Зачтено</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td data-label="Дата">23.08.2026</td>
              <td data-label="Сумма" class="num">2&nbsp;214,00&nbsp;₽</td>
              <td data-label="Назначение">НДФЛ → ФНС (ЕНП)</td>
              <td data-label="Период">август 2026 (часть)</td>
              <td data-label="Зачтено"><span class="badge badge--ok">уже зачтено</span></td>
            </tr>
            <tr>
              <td data-label="Дата">23.09.2026</td>
              <td data-label="Сумма" class="num">1&nbsp;137,00&nbsp;₽</td>
              <td data-label="Назначение">НДФЛ → ФНС (ЕНП)</td>
              <td data-label="Период">август 2026 (доплата)</td>
              <td data-label="Зачтено"><span class="badge badge--pending">еще не зачтено</span></td>
            </tr>
            <tr>
              <td data-label="Дата">23.09.2026</td>
              <td data-label="Сумма" class="num">323,00&nbsp;₽</td>
              <td data-label="Назначение">НДФЛ → ФНС (ЕНП)</td>
              <td data-label="Период">сентябрь 2026 (аванс)</td>
              <td data-label="Зачтено"><span class="badge badge--pending">еще не зачтено</span></td>
            </tr>
            <tr>
              <td data-label="Дата">23.09.2026</td>
              <td data-label="Сумма" class="num">7&nbsp;733,16&nbsp;₽</td>
              <td data-label="Назначение">Взносы 30% → ФНС (ЕНП)</td>
              <td data-label="Период">август 2026</td>
              <td data-label="Зачтено"><span class="badge badge--pending">еще не зачтено</span></td>
            </tr>
          </tbody>
        </table>
      </div>

      <div class="po-notes">
        <div class="po-note">
          <p class="po-note__label">Август</p>
          <p class="po-note__text">
            НДФЛ <strong>3&nbsp;351,00&nbsp;₽</strong> (2&nbsp;214,00 + 1&nbsp;137,00), взносы
            <strong>7&nbsp;733,16&nbsp;₽</strong> — соответствуют начислению 25&nbsp;777,20&nbsp;₽.
          </p>
        </div>
        <div class="po-note">
          <p class="po-note__label">Сентябрь</p>
          <p class="po-note__text">
            НДФЛ с аванса <strong>323,00&nbsp;₽</strong> уплачен; доплата ≈ 323,00&nbsp;₽ и взносы
            1&nbsp;489,88&nbsp;₽ — срок до 28.10.2026.
          </p>
        </div>
        <div class="po-note">
          <p class="po-note__label">Травматизм 0,2% СФР</p>
          <p class="po-note__text">
            Август ≈ 51,55&nbsp;₽ / сентябрь 9,93&nbsp;₽ — в данной выписке не подтверждены, проверить отдельно.
          </p>
        </div>
      </div>
    </section>

    <section id="contract" class="po-block">
      <div class="po-block__head">
        <span class="po-num">03</span>
        <div>
          <h2 class="po-h2">Документы</h2>
          <p class="po-sub">Трудовой договор от 06.08.2026 и допсоглашение об 0,125 ставки</p>
        </div>
      </div>
      <p class="po-lead">
        Канонические файлы для кадрового учёта и бухгалтерии. При расхождении версий приоритет у PDF
        с подписями сторон.
      </p>
      <div class="doc-actions">
        <a class="btn btn--primary" href="трудовой-договор/трудовой-договор-Ходжиматов-бессрочный-06.08.2026.pdf" download>Трудовой договор (PDF)</a>
        <a class="btn btn--ghost" href="трудовой-договор/трудовой-договор-Ходжиматов-бессрочный-06.08.2026.docx" download>Трудовой договор (Word)</a>
        <a class="btn btn--ghost" href="трудовой-договор/допсоглашение-0.125-ставки-Ходжиматов.pdf" download>Допсоглашение 0,125 (PDF)</a>
      </div>
    </section>

    <section id="registry" class="po-block">
      <div class="po-block__head">
        <span class="po-num">04</span>
        <div>
          <h2 class="po-h2">Реестр личных покупок из Сбермаркета</h2>
          <p class="po-sub">К перевыставлению Фуркатжону · итого <strong>9&nbsp;853,56&nbsp;₽</strong> · 40 позиций</p>
        </div>
      </div>
      <p class="po-lead">
        Согласованные личные позиции из заказов доставки Сбермаркет.
        Товары кофейни, цитрусы и сервисные сборы не учитываются.
      </p>
      <div class="doc-actions" style="margin-bottom:0.85rem">
        <a class="btn btn--ghost" href="заказы-купер/реестр-перевыставление-Фурик.xlsx" download>Скачать реестр (Excel)</a>
      </div>
      <div class="table-wrap">
        <table class="table-stack table-registry">
          <thead>
            <tr>
              <th>Дата</th>
              <th>№ заказа</th>
              <th>Товар</th>
              <th class="num">Сумма</th>
              <th>Магазин</th>
            </tr>
          </thead>
          <tbody>
{rows}
            <tr class="row-total">
              <td data-label="Итого" colspan="3"><strong>Итого к перевыставлению</strong></td>
              <td data-label="Сумма" class="num"><strong>9&nbsp;853,56&nbsp;₽</strong></td>
              <td data-label="Позиции">40 позиций</td>
            </tr>
          </tbody>
        </table>
      </div>
    </section>

    <p class="page-foot">
      Сопутствующий материал: <a href="index.html">инструкция по приёму сотрудника</a>.
      Суммы по оплате труда — расчётный ориентир; перед платежами рекомендуется сверка с бухгалтером.
    </p>
  </main>

  <script>
    (function () {{
      const toggle = document.getElementById("navToggle");
      const nav = document.getElementById("topnav");
      if (!toggle || !nav) return;
      toggle.addEventListener("click", () => {{
        const open = nav.classList.toggle("is-open");
        toggle.setAttribute("aria-expanded", open ? "true" : "false");
      }});
      nav.querySelectorAll("a").forEach((a) => {{
        a.addEventListener("click", () => {{
          nav.classList.remove("is-open");
          toggle.setAttribute("aria-expanded", "false");
        }});
      }});
    }})();
  </script>
</body>
</html>
"""

(DOCS / "raschety.html").write_text(html, encoding="utf-8")
print(f"ok rows={n} total={total:.2f} bytes={len(html)}")
