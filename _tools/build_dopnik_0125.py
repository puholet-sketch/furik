# -*- coding: utf-8 -*-
"""Допсоглашение 0,125 ставки (1 ч/день) — гибкое окно 08:00–18:00."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"D:\projects\Фурик")
OUT = ROOT / "для-бухгалтерии" / "допсоглашение-0.125-ставки-Ходжиматов.docx"
DOCS = ROOT / "docs" / "трудовой-договор"

INK = RGBColor(0x1A, 0x1A, 0x1A)
TEXT = RGBColor(0x33, 0x33, 0x33)

OKLAD_FULL = 39730
RATE = 0.125
OKLAD = round(OKLAD_FULL * RATE, 2)  # 4966.25
AGREEMENT_DATE = "25.08.2026"
AGREEMENT_DATE_LONG = "«25» августа 2026 г."

E = {
    "fio": "Сорванова Анна Александровна",
    "fio_short": "Сорванова А. А.",
    "ogrnip": "322774600583080",
    "inn": "772973703990",
    "address": "г. Москва, ул. Василия Ланового, д. 3, кв. 243",
    "bank": 'МОСКОВСКИЙ ФИЛИАЛ АО КБ "МОДУЛЬБАНК"',
    "bik": "044525092",
    "ks": "30101810645250000092",
    "rs": "40802810370010393644",
    "email": "SORVANOVAAA@GMAIL.COM",
}
W = {
    "fio": "Ходжиматов Фуркатжон Махамаджонович",
    "fio_short": "Ходжиматов Ф. М.",
    "passport": "FA6253664",
    "passport_until": "28.08.2027",
    "address": "г. Люберцы, ул. 8 Марта, д. 53, кв. 47",
    "mig_series": "45 26",
    "mig_number": "0767385",
    "patent": "серия 77 № 2600347821",
    "patent_profession": "Бармен",
    "snils": "229-173-224 63",
    "inn": "772430430429",
    "phone": "8 (901) 797-57-53",
}
B = {
    "account": "40820810338110973759",
    "bank": "ПАО Сбербанк",
    "bik": "044525225",
}


def set_run(run, *, size=11, bold=False, color=TEXT):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def set_nil_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    for child in list(tbl_pr):
        if child.tag == qn("w:tblBorders"):
            tbl_pr.remove(child)
    tbl_pr.append(borders)
    if tbl.tblPr is None:
        tbl.insert(0, tbl_pr)


def set_table_borders(table, color="BBBBBB", sz="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    for child in list(tbl_pr):
        if child.tag == qn("w:tblBorders"):
            tbl_pr.remove(child)
    tbl_pr.append(borders)


def set_cell_margins(cell, top=30, bottom=30, left=40, right=40):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for m, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    for child in list(tc_pr):
        if child.tag == qn("w:tcMar"):
            tc_pr.remove(child)
    tc_pr.append(tc_mar)


def set_cell_shading(cell, fill_hex: str):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def cell_para(cell, text, *, bold=False, size=10, center=False, before=0, after=4):
    if len(cell.paragraphs) == 1 and not cell.paragraphs[0].text.strip():
        p = cell.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
    else:
        p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=INK if bold else TEXT)


def add_kv_table(parent, rows, label_width_cm=2.2, value_width_cm=5.8, font_size=8):
    table = parent.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders(table)
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.width = Cm(label_width_cm)
        c1.width = Cm(value_width_cm)
        set_cell_margins(c0)
        set_cell_margins(c1)
        set_cell_shading(c0, "FFFFFF")
        set_cell_shading(c1, "FFFFFF")
        p0, p1 = c0.paragraphs[0], c1.paragraphs[0]
        for r in list(p0.runs):
            r._element.getparent().remove(r._element)
        for r in list(p1.runs):
            r._element.getparent().remove(r._element)
        p0.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        set_run(r0, size=font_size, bold=True, color=INK)
        r1 = p1.add_run(v)
        set_run(r1, size=font_size, color=TEXT)


def add_p(doc, text="", *, bold=False, size=11, center=False, color=TEXT):
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, color=color)
    return p


def main():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = TEXT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15

    add_p(doc, "ДОПОЛНИТЕЛЬНОЕ СОГЛАШЕНИЕ № 1", bold=True, size=14, center=True, color=INK)
    add_p(doc, "к трудовому договору от 06.08.2026", bold=True, size=12, center=True, color=INK)
    add_p(doc, "с работником — иностранным гражданином", center=True, size=11, color=INK)
    add_p(
        doc,
        f"г. Москва                                                    {AGREEMENT_DATE_LONG}",
        center=True,
    )

    add_p(
        doc,
        f"Индивидуальный предприниматель {E['fio']}, ОГРНИП {E['ogrnip']}, "
        f"ИНН {E['inn']}, зарегистрированная по адресу: {E['address']}, "
        f"именуемая в дальнейшем «Работодатель», с одной стороны, и "
        f"гражданин Республики Узбекистан {W['fio']}, "
        f"именуемый в дальнейшем «Работник», с другой стороны, совместно именуемые «Стороны», "
        f"заключили настоящее дополнительное соглашение к трудовому договору от 06.08.2026 "
        f"(далее — Договор) о нижеследующем.",
    )

    add_p(doc, "1. Предмет соглашения", bold=True, color=INK)
    add_p(
        doc,
        "1.1. Стороны в соответствии со статьёй 93 Трудового кодекса Российской Федерации "
        "устанавливают Работнику неполное рабочее время — работу на условиях "
        "0,125 (одной восьмой) ставки (0,125 полной нормы рабочего времени).",
    )
    add_p(
        doc,
        "1.2. Настоящее дополнительное соглашение изменяет условия Договора в части "
        "режима рабочего времени, нормы рабочего времени и оплаты труда в редакции, "
        "изложенной ниже. Остальные условия Договора сохраняют силу в неизменном виде, "
        "поскольку не противоречат настоящему соглашению.",
    )

    add_p(doc, "2. Режим рабочего времени", bold=True, color=INK)
    add_p(
        doc,
        "2.1. Работнику устанавливается пятидневная рабочая неделя с двумя выходными днями "
        "(суббота и воскресенье).",
    )
    add_p(
        doc,
        "2.2. Продолжительность рабочего времени — 5 (пять) часов в неделю "
        "(1 (один) час в день), что соответствует 0,125 ставки от нормальной "
        "продолжительности рабочего времени 40 часов в неделю.",
    )
    add_p(
        doc,
        "2.3. Ежедневная работа продолжительностью 1 (один) час выполняется в пределах "
        "временного окна с 08:00 до 18:00 по московскому времени в месте работы. "
        "Конкретное время начала и окончания указанного часа (например, с 10:00 до 11:00 "
        "либо иной непрерывный часовой интервал внутри окна) не является жёстко "
        "зафиксированным в настоящем соглашении и определяется по согласованию Сторон "
        "(устно, посредством сообщений или иным способом, позволяющим подтвердить "
        "договорённость) применительно к каждому рабочему дню либо на период, "
        "согласованный Сторонами, с учётом производственной необходимости Работодателя "
        "и возможности Работника. При отсутствии иного согласования на конкретный день "
        "Работник обязан явиться для исполнения трудовой функции в интервале, "
        "заранее сообщённом Работодателем в пределах окна 08:00–18:00.",
    )
    add_p(
        doc,
        "2.4. Перерыв для отдыха и питания при продолжительности ежедневной работы "
        "1 (один) час не устанавливается.",
    )
    add_p(doc, "2.5. Место работы без изменений: г. Москва, ул. Киевская, д. 7, к. 2.")
    add_p(
        doc,
        "2.6. Должность (трудовая функция) без изменений: Бармен; работа осуществляется "
        "в пределах профессии, указанной в патенте Работника, и на территории г. Москвы.",
    )
    add_p(
        doc,
        "2.7. Учёт рабочего времени ведётся Работодателем исходя из фактически "
        "отработанного часа в соответствующий рабочий день в пределах согласованного "
        "интервала внутри окна 08:00–18:00.",
    )

    add_p(doc, "3. Оплата труда", bold=True, color=INK)
    add_p(
        doc,
        f"3.1. С даты вступления в силу настоящего соглашения Работнику устанавливается "
        f"должностной оклад за работу на условиях 0,125 ставки в размере "
        f"{OKLAD:.2f} руб. ({OKLAD:.2f} рублей) в месяц "
        f"(пропорционально 0,125 от оклада {OKLAD_FULL} руб. за полную норму рабочего времени).",
    )
    add_p(
        doc,
        "3.2. При неполном отработанном месяце оплата производится пропорционально "
        "отработанному времени.",
    )
    add_p(
        doc,
        "3.3. Заработная плата выплачивается путём перечисления на банковский счёт Работника "
        "в порядке и сроки, установленные Договором (безналично). Выплата наличными не допускается.",
    )
    add_p(
        doc,
        "3.4. Оплата труда за полностью отработанную норму рабочего времени, установленную "
        "настоящим соглашением, не может быть ниже величины, исчисленной пропорционально "
        "минимальной заработной плате в г. Москве (с 01.01.2026 — 39 730 руб. за полную норму).",
    )

    add_p(doc, "4. Срок действия", bold=True, color=INK)
    add_p(
        doc,
        "4.1. Настоящее дополнительное соглашение вступает в силу с 25.08.2026 "
        "и действует в течение срока действия Договора, если Стороны не изменят условия иным "
        "письменным соглашением.",
    )
    add_p(
        doc,
        "4.2. Соглашение составлено в двух экземплярах, имеющих одинаковую юридическую силу, "
        "по одному для каждой Стороны.",
    )

    add_p(doc, "5. Реквизиты и подписи сторон", bold=True, color=INK)

    outer = doc.add_table(rows=1, cols=2)
    set_nil_borders(outer)
    left, right = outer.rows[0].cells
    left.width = Cm(8.6)
    right.width = Cm(8.6)
    set_cell_margins(left, 40, 40, 20, 80)
    set_cell_margins(right, 40, 40, 80, 20)

    cell_para(left, "РАБОТОДАТЕЛЬ", bold=True, size=11, center=True, after=6)
    add_kv_table(
        left,
        [
            ("Статус", f"ИП {E['fio']}"),
            ("ОГРНИП", E["ogrnip"]),
            ("ИНН", E["inn"]),
            ("Адрес", E["address"]),
            ("р/с", E["rs"]),
            ("Банк", E["bank"]),
            ("БИК", E["bik"]),
            ("к/с", E["ks"]),
            ("E-mail", E["email"]),
        ],
        label_width_cm=2.2,
        value_width_cm=5.8,
        font_size=8,
    )
    cell_para(left, f"_________________ / {E['fio_short']} /", size=8, before=10, after=2)
    cell_para(
        left,
        f"М.П. (при наличии)     Дата: {AGREEMENT_DATE}",
        size=8,
        after=0,
    )

    cell_para(right, "РАБОТНИК", bold=True, size=11, center=True, after=6)
    add_kv_table(
        right,
        [
            ("ФИО", W["fio"]),
            ("Гражданство", "Республика Узбекистан"),
            ("Паспорт", f"№ {W['passport']}, до {W['passport_until']}"),
            ("Адрес", W["address"]),
            ("Мигр. карта", f"серия {W['mig_series']} № {W['mig_number']}"),
            ("Патент", f"{W['patent']}, {W['patent_profession']}"),
            ("СНИЛС", W["snils"]),
            ("ИНН", W["inn"]),
            ("Телефон", W["phone"]),
            ("Счёт (безнал)", B["account"]),
            ("Банк", f"{B['bank']}, БИК {B['bik']}"),
        ],
        label_width_cm=2.6,
        value_width_cm=5.4,
        font_size=8,
    )
    cell_para(right, "Экземпляр соглашения получил:", size=8, before=10, after=2)
    cell_para(right, f"_________________ / {W['fio_short']} /", size=8, after=2)
    cell_para(right, f"Дата: {AGREEMENT_DATE}", size=8, after=0)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    DOCS.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT, DOCS / OUT.name)
    print("WROTE", OUT)
    print("OKLAD", OKLAD)


if __name__ == "__main__":
    main()
