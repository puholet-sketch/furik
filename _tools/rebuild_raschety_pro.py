# -*- coding: utf-8 -*-
"""Rebuild professional raschety.html from Excel registry."""
from pathlib import Path
from html import escape
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
XLSX = ROOT / "заказы-купер" / "реестр-перевыставление-Фурик.xlsx"


def rub(x: float) -> str:
    return f"{x:,.2f}".replace(",", "\u00a0").replace(".", ",")


def fmt_date(d: str) -> str:
    # 2026-01-29 -> 29.01.2026
    y, m, day = str(d)[:10].split("-")
    return f"{day}.{m}.{y}"


wb = load_workbook(XLSX)
ws = wb.active
trs = []
total = 0.0
n = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    d, no, pos, qty, unit, sm, src = r[:7]
    sm = float(sm)
    total += sm
    n += 1
    store = (src or "").replace(" / Купер", "").strip()
    trs.append(
        "<tr>"
        f'<td data-label="Дата">{escape(fmt_date(d))}</td>'
        f'<td data-label="№ заказа"><code>{escape(str(no or ""))}</code></td>'
        f'<td data-label="Товар">{escape(pos or "")}</td>'
        f'<td data-label="Сумма" class="num">{rub(sm)}&nbsp;₽</td>'
        f'<td data-label="Магазин">{escape(store)}</td>'
        "</tr>"
    )

assert n == 40, n
assert abs(total - 9853.56) < 0.02, total
rows = "\n".join(trs)

html = f"""<!DOCTYPE html>
<html lang="ru">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1" />
  <title>Расчёты и взаиморасчёты — Ходжиматов Фуркатжон</title>
  <link rel="stylesheet" href="styles.css" />
  <style>
    .table-registry {{ min-width: 720px; }}
    .table-registry col.c-date {{ width: 11%; }}
    .table-registry col.c-order {{ width: 16%; }}
    .table-registry col.c-item {{ width: 43%; }}
    .table-registry col.c-sum {{ width: 14%; }}
    .table-registry col.c-store {{ width: 16%; }}
    .table-registry code {{
      font-size: 0.82em;
      background: var(--blue-soft);
      padding: 0.1em 0.35em;
      border-radius: 4px;
      white-space: nowrap;
    }}
    .section-intro {{
      margin: 0 0 14px;
      color: var(--muted);
      font-size: 0.95rem;
      max-width: 52rem;
    }}
    .split-actions {{
      display: flex;
      flex-wrap: wrap;
      gap: 10px;
      align-items: center;
    }}
    .note {{
      margin-top: 14px;
      font-size: 0.92rem;
      color: var(--muted);
    }}
    .page-foot {{
      padding: 8px 16px 40px;
      max-width: var(--max);
      margin: 0 auto;
      font-size: 0.9rem;
      color: var(--muted);
    }}
  </style>
</head>
<body>
  <header class="topbar" id="topbar">
    <div class="topbar__inner">
      <div class="brand">
        <span class="brand__mark" aria-hidden="true"></span>
        <div>
          <div class="brand__name">Взаиморасчёты</div>
          <div class="brand__sub">Ходжиматов Фуркатжон · 2026</div>
        </div>
      </div>
      <button class="nav-toggle" type="button" aria-label="Меню" aria-expanded="false" id="navToggle">
        <span></span>
      </button>
      <nav class="topnav" aria-label="Разделы" id="topnav">
        <a href="index.html">← Инструкция</a>
        <a href="#zp">1. Оплата труда</a>
        <a href="#paid">2. Бюджет</a>
        <a href="#contract">3. Документы</a>
        <a href="#registry">4. Покупки</a>
      </nav>
    </div>
  </header>

  <main>
    <section class="hero block is-visible">
      <p class="eyebrow">ИП Сорванова · служебная сводка · обновлено 27.09.2026</p>
      <h1>Расчёты по сотруднику Ходжиматову Фуркатжону</h1>
      <p class="lead">
        Сводка по официальной оплате труда (0,125 ставки с 25.08.2026), платежам в бюджет
        по банковской выписке и реестру личных покупок для перевыставления.
      </p>
      <div class="hero__actions">
        <a class="btn btn--primary" href="#zp">К разделу 1</a>
        <a class="btn btn--ghost" href="#registry">К реестру</a>
        <a class="btn btn--pdf" href="трудовой-договор/трудовой-договор-Ходжиматов-бессрочный-06.08.2026.pdf" download>ТД PDF</a>
      </div>
    </section>

    <div class="facts">
      <div class="fact">
        <span class="fact__label">На руки / мес.</span>
        <span class="fact__value">4&nbsp;320,25&nbsp;₽</span>
        <span class="fact__hint">оклад 4&nbsp;966,25 − НДФЛ</span>
      </div>
      <div class="fact">
        <span class="fact__label">Расход ИП / мес.</span>
        <span class="fact__value">6&nbsp;466,06&nbsp;₽</span>
        <span class="fact__hint">оклад + взносы + травматизм</span>
      </div>
      <div class="fact">
        <span class="fact__label">Налоги и взносы</span>
        <span class="fact__value">2&nbsp;145,81&nbsp;₽</span>
        <span class="fact__hint">НДФЛ + 30% + 0,2%</span>
      </div>
      <div class="fact">
        <span class="fact__label">Реестр покупок</span>
        <span class="fact__value">9&nbsp;853,56&nbsp;₽</span>
        <span class="fact__hint">40 позиций · 2026</span>
      </div>
    </div>

    <section id="zp" class="block">
      <div class="step-head">
        <span class="step-num">1</span>
        <div>
          <h2>Официальная оплата труда (0,125 ставки)</h2>
          <p class="muted">Ходжиматов Фуркатжон Махамаджонович · допсоглашение с 25.08.2026</p>
        </div>
      </div>
      <p class="section-intro">
        Официальная часть оплаты труда — 1/8 полной ставки (оклад МЗП Москвы 39&nbsp;730&nbsp;₽).
        Ниже — структура начисления, удержаний и обязательных платежей работодателя за полный месяц.
      </p>

      <div class="table-wrap">
        <table class="table-stack">
          <thead>
            <tr>
              <th>Параметр</th>
              <th class="num">Полная ставка</th>
              <th class="num">0,125 ставки</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td data-label="Параметр">Оклад (начисление)</td>
              <td data-label="Полная ставка" class="num">39&nbsp;730,00&nbsp;₽</td>
              <td data-label="0,125 ставки" class="num"><strong>4&nbsp;966,25&nbsp;₽</strong></td>
            </tr>
            <tr>
              <td data-label="Параметр">Норма рабочего времени</td>
              <td data-label="Полная ставка" class="num">40 ч / нед.</td>
              <td data-label="0,125 ставки" class="num">5 ч / нед. (1 ч / день)</td>
            </tr>
          </tbody>
        </table>
      </div>

      <div class="table-wrap" style="margin-top:16px">
        <table class="table-stack">
          <thead>
            <tr>
              <th>Статья</th>
              <th class="num">Сумма</th>
              <th>Куда / кто</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td data-label="Статья">Выплата сотруднику</td>
              <td data-label="Сумма" class="num"><strong>4&nbsp;320,25&nbsp;₽</strong></td>
              <td data-label="Куда">на счёт в ПАО Сбербанк</td>
            </tr>
            <tr>
              <td data-label="Статья">НДФЛ 13%</td>
              <td data-label="Сумма" class="num">646,00&nbsp;₽</td>
              <td data-label="Куда">удержание из оклада → ФНС (ЕНП)</td>
            </tr>
            <tr>
              <td data-label="Статья">Страховые взносы 30%</td>
              <td data-label="Сумма" class="num">1&nbsp;489,88&nbsp;₽</td>
              <td data-label="Куда">за счёт ИП → ФНС (ЕНП)</td>
            </tr>
            <tr>
              <td data-label="Статья">Взносы на травматизм 0,2%</td>
              <td data-label="Сумма" class="num">9,93&nbsp;₽</td>
              <td data-label="Куда">за счёт ИП → СФР</td>
            </tr>
            <tr class="row-total">
              <td data-label="Статья">Итого НДФЛ + взносы + травматизм</td>
              <td data-label="Сумма" class="num"><strong>2&nbsp;145,81&nbsp;₽</strong></td>
              <td data-label="Куда">в месяц при 0,125 ставки</td>
            </tr>
          </tbody>
        </table>
      </div>

      <div class="alert alert--info" style="margin-top:14px">
        Совокупный расход ИП по официальной части:
        4&nbsp;966,25&nbsp;₽ (оклад) + 1&nbsp;489,88&nbsp;₽ + 9,93&nbsp;₽ =
        <strong>6&nbsp;466,06&nbsp;₽</strong> в месяц. НДФЛ входит в оклад и сверху не добавляется.
        Суммы сверх официальной части учитываются отдельно в рамках утверждённого недельного бюджета проекта
        и в настоящую таблицу не включены.
      </div>
    </section>

    <section id="paid" class="block">
      <div class="step-head">
        <span class="step-num">2</span>
        <div>
          <h2>Уплачено в бюджет (по выписке)</h2>
          <p class="muted">Подтверждённые платежи по Фуркатжону · август–сентябрь 2026</p>
        </div>
      </div>
      <p class="section-intro">
        Строки ниже сверены с банковской выпиской. Указаны только фактически проведённые платежи;
        плановые суммы с ненаступившим сроком отмечены в примечании.
      </p>

      <div class="table-wrap">
        <table class="table-stack">
          <thead>
            <tr>
              <th>Дата</th>
              <th class="num">Сумма</th>
              <th>Назначение</th>
              <th>Период</th>
              <th>Зачтено</th>
            </tr>
          </thead>
          <tbody>
            <tr>
              <td data-label="Дата">23.08.2026</td>
              <td data-label="Сумма" class="num">2&nbsp;214,00&nbsp;₽</td>
              <td data-label="Назначение">НДФЛ → ФНС (ЕНП)</td>
              <td data-label="Период">август 2026 (часть)</td>
              <td data-label="Зачтено">уже зачтено</td>
            </tr>
            <tr>
              <td data-label="Дата">23.09.2026</td>
              <td data-label="Сумма" class="num">1&nbsp;137,00&nbsp;₽</td>
              <td data-label="Назначение">НДФЛ → ФНС (ЕНП)</td>
              <td data-label="Период">август 2026 (доплата)</td>
              <td data-label="Зачтено">еще не зачтено</td>
            </tr>
            <tr>
              <td data-label="Дата">23.09.2026</td>
              <td data-label="Сумма" class="num">323,00&nbsp;₽</td>
              <td data-label="Назначение">НДФЛ → ФНС (ЕНП)</td>
              <td data-label="Период">сентябрь 2026 (аванс)</td>
              <td data-label="Зачтено">еще не зачтено</td>
            </tr>
            <tr>
              <td data-label="Дата">23.09.2026</td>
              <td data-label="Сумма" class="num">7&nbsp;733,16&nbsp;₽</td>
              <td data-label="Назначение">Взносы 30% → ФНС (ЕНП)</td>
              <td data-label="Период">август 2026</td>
              <td data-label="Зачтено">еще не зачтено</td>
            </tr>
          </tbody>
        </table>
      </div>

      <p class="note">
        Август: НДФЛ итого <strong>3&nbsp;351,00&nbsp;₽</strong> (2&nbsp;214,00 + 1&nbsp;137,00) и взносы
        <strong>7&nbsp;733,16&nbsp;₽</strong> — соответствуют начислению 25&nbsp;777,20&nbsp;₽.
        Сентябрь: НДФЛ с аванса <strong>323,00&nbsp;₽</strong> уплачен; доплата НДФЛ ≈ 323,00&nbsp;₽ и взносы
        1&nbsp;489,88&nbsp;₽ — срок до 28.10.2026.
        Платежи на травматизм 0,2% в СФР (август ≈ 51,55&nbsp;₽, сентябрь 9,93&nbsp;₽) в данной выписке
        не подтверждены и требуют отдельной проверки.
      </p>
    </section>

    <section id="contract" class="block">
      <div class="step-head">
        <span class="step-num">3</span>
        <div>
          <h2>Документы</h2>
          <p class="muted">Трудовой договор от 06.08.2026 и допсоглашение об 0,125 ставки</p>
        </div>
      </div>
      <p class="section-intro">
        Канонические файлы для кадрового учёта и бухгалтерии. При расхождении версий приоритет у PDF
        с подписями сторон.
      </p>
      <div class="split-actions">
        <a class="btn btn--pdf" href="трудовой-договор/трудовой-договор-Ходжиматов-бессрочный-06.08.2026.pdf" download>Трудовой договор (PDF)</a>
        <a class="btn btn--ghost" href="трудовой-договор/трудовой-договор-Ходжиматов-бессрочный-06.08.2026.docx" download>Трудовой договор (Word)</a>
        <a class="btn btn--ghost" href="трудовой-договор/допсоглашение-0.125-ставки-Ходжиматов.pdf" download>Допсоглашение 0,125 (PDF)</a>
      </div>
    </section>

    <section id="registry" class="block">
      <div class="step-head">
        <span class="step-num">4</span>
        <div>
          <h2>Реестр личных покупок из Сбермаркета</h2>
          <p class="muted">К перевыставлению Фуркатжону · итого <strong>9&nbsp;853,56&nbsp;₽</strong> · 40 позиций</p>
        </div>
      </div>
      <p class="section-intro">
        В реестр включены согласованные личные позиции из заказов доставки Сбермаркет.
        Товары кофейни, цитрусы и сервисные сборы не учитываются. Файл Excel можно скачать ниже.
      </p>
      <div class="split-actions" style="margin-bottom:12px">
        <a class="btn btn--primary" href="заказы-купер/реестр-перевыставление-Фурик.xlsx" download>Скачать реестр (Excel)</a>
      </div>
      <div class="table-wrap">
        <table class="table-stack table-registry">
          <colgroup>
            <col class="c-date" /><col class="c-order" /><col class="c-item" /><col class="c-sum" /><col class="c-store" />
          </colgroup>
          <thead>
            <tr>
              <th>Дата</th>
              <th>№ заказа</th>
              <th>Товар</th>
              <th class="num">Сумма</th>
              <th>Магазин</th>
            </tr>
          </thead>
          <tbody>
{rows}
            <tr class="row-total">
              <td data-label="Итого" colspan="3"><strong>Итого к перевыставлению</strong></td>
              <td data-label="Сумма" class="num"><strong>9&nbsp;853,56&nbsp;₽</strong></td>
              <td data-label="Позиции">40 позиций</td>
            </tr>
          </tbody>
        </table>
      </div>
    </section>

    <p class="page-foot">
      Сопутствующий материал: <a href="index.html">инструкция по приёму сотрудника</a>.
      Суммы по оплате труда — расчётный ориентир; перед платежами рекомендуется сверка с бухгалтером.
    </p>
  </main>

  <script>
    (function () {{
      const toggle = document.getElementById("navToggle");
      const nav = document.getElementById("topnav");
      if (!toggle || !nav) return;
      toggle.addEventListener("click", () => {{
        const open = nav.classList.toggle("is-open");
        toggle.setAttribute("aria-expanded", open ? "true" : "false");
      }});
      nav.querySelectorAll("a").forEach((a) => {{
        a.addEventListener("click", () => {{
          nav.classList.remove("is-open");
          toggle.setAttribute("aria-expanded", "false");
        }});
      }});
    }})();
  </script>
</body>
</html>
"""

(DOCS / "raschety.html").write_text(html, encoding="utf-8")
print(f"ok rows={n} total={total:.2f} bytes={len(html)}")
