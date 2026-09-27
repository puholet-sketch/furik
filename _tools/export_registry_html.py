# -*- coding: utf-8 -*-
from pathlib import Path
from html import escape
from openpyxl import load_workbook
from shutil import copy2
from datetime import date, datetime

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "заказы-купер" / "реестр-перевыставление-Фурик.xlsx"
DOCS = ROOT / "docs"
(DOCS / "заказы-купер").mkdir(parents=True, exist_ok=True)
copy2(XLSX, DOCS / "заказы-купер" / "реестр-перевыставление-Фурик.xlsx")


def rub(x: float) -> str:
    return f"{x:,.2f}".replace(",", "\u00a0").replace(".", ",")


def fmt_date(v) -> str:
    if isinstance(v, datetime):
        return v.strftime("%d.%m.%Y")
    if isinstance(v, date):
        return v.strftime("%d.%m.%Y")
    s = str(v or "")[:10]
    if len(s) == 10 and s[4] == "-":
        y, m, d = s.split("-")
        return f"{d}.{m}.{y}"
    return s


wb = load_workbook(XLSX, data_only=True)
ws = wb.active

header_row = 1
for r in range(1, min(ws.max_row, 40) + 1):
    v = ws.cell(r, 1).value
    if v and str(v).strip().lower().startswith("дата"):
        header_row = r
        break

rows = []
for r in range(header_row + 1, ws.max_row + 1):
    vals = [ws.cell(r, c).value for c in range(1, 8)]
    d, no, pos, qty, unit, sm, src = vals
    if not d:
        continue
    if pos and "итого" in str(pos).lower():
        continue
    rows.append(
        (
            fmt_date(d),
            no or "",
            pos or "",
            float(sm or 0),
            (src or "").replace(" / Купер", "").strip(),
        )
    )

total = sum(r[3] for r in rows)
trs = []
for d, no, pos, sm, store in rows:
    trs.append(
        "<tr>"
        f'<td data-label="Дата">{escape(d)}</td>'
        f'<td data-label="№ заказа"><code>{escape(str(no))}</code></td>'
        f'<td data-label="Товар">{escape(pos)}</td>'
        f'<td data-label="Сумма" class="num">{rub(sm)}\u00a0₽</td>'
        f'<td data-label="Магазин">{escape(store)}</td>'
        "</tr>"
    )

fragment = "\n".join(trs)
meta = f"{len(rows)}|{rub(total)}"
(DOCS / "_reg_fragment.html").write_text(fragment, encoding="utf-8")
(DOCS / "_reg_meta.txt").write_text(meta, encoding="utf-8")
print(meta)
