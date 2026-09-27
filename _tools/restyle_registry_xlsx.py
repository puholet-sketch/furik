# -*- coding: utf-8 -*-
"""Переоформление реестра Сбермаркета в пастельном стиле REMONT (THEME_PASTEL)."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "заказы-купер" / "реестр-перевыставление-Фурик.xlsx"
OUTS = [
    SRC,
    ROOT / "docs" / "заказы-купер" / "реестр-перевыставление-Фурик.xlsx",
]

# Канон REMONT: build_photo_registry.py / THEME_PASTEL в generate_stage_contract.py
HEADER_SHADE = "FFF2C2"
KEY_SHADE = "FFF8E7"
ZEBRA_SHADE = "FFFFFF"
TOTAL_SHADE = "FFF0C8"
BORDER_SOFT = "E5DDD2"
ACCENT = "D4B552"
INK = "000000"

FONT_HEAD = "Arial Black"
FONT_BODY = "Arial"

fill_header = PatternFill("solid", fgColor=HEADER_SHADE)
fill_key = PatternFill("solid", fgColor=KEY_SHADE)
fill_zebra = PatternFill("solid", fgColor=ZEBRA_SHADE)
fill_total = PatternFill("solid", fgColor=TOTAL_SHADE)

thin_soft = Side(style="thin", color=BORDER_SOFT)
med_accent = Side(style="medium", color=ACCENT)
border_soft = Border(left=thin_soft, right=thin_soft, top=thin_soft, bottom=thin_soft)
border_accent = Border(left=med_accent, right=med_accent, top=med_accent, bottom=med_accent)

font_title = Font(name=FONT_HEAD, size=16, bold=True, color=INK)
font_sub = Font(name=FONT_BODY, size=10, color=INK)
font_head = Font(name=FONT_BODY, size=10, bold=True, color=INK)
font_body = Font(name=FONT_BODY, size=9, color=INK)
font_body_bold = Font(name=FONT_BODY, size=9, bold=True, color=INK)

align_wrap = Alignment(wrap_text=True, vertical="center")
align_center = Alignment(wrap_text=True, vertical="center", horizontal="center")
align_left = Alignment(wrap_text=True, vertical="center", horizontal="left")
align_right = Alignment(wrap_text=True, vertical="center", horizontal="right")

HEADERS = [
    "Дата заказа",
    "№ заказа",
    "Позиция",
    "Кол-во",
    "Стоимость",
    "Сумма",
    "Источник",
]
NCOLS = len(HEADERS)
MONEY_FMT = '#,##0.00'


def put(ws, r, c, value, *, fill=None, font=None, border=None, align=None, number_format=None):
    cell = ws.cell(r, c, value)
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if border is not None:
        cell.border = border
    if align is not None:
        cell.alignment = align
    if number_format is not None:
        cell.number_format = number_format
    return cell


def parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v.strip():
        s = v.strip()[:10]
        try:
            return date.fromisoformat(s)
        except ValueError:
            return s
    return v


def load_data(path: Path) -> tuple[list[tuple], list[str], float]:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows: list[tuple] = []
    total = 0.0
    # Find header row by first cell
    header_row = None
    for r in range(1, min(ws.max_row, 30) + 1):
        v = ws.cell(r, 1).value
        if v and str(v).strip().lower().startswith("дата"):
            header_row = r
            break
    if header_row is None:
        header_row = 1

    for r in range(header_row + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, 8)]
        d, no, pos, qty, unit, sm, src = vals
        if not d and pos and "итого" in str(pos).lower():
            total = float(sm or 0)
            continue
        if not d:
            continue
        rows.append(
            (
                parse_date(d),
                str(no or ""),
                str(pos or ""),
                float(qty or 0),
                float(unit or 0),
                float(sm or 0),
                str(src or ""),
            )
        )

    if not total:
        total = round(sum(r[5] for r in rows), 2)

    rules: list[str] = []
    if "Правила" in wb.sheetnames:
        for r in wb["Правила"].iter_rows(min_row=1, max_col=1, values_only=True):
            if r[0]:
                rules.append(str(r[0]))
    return rows, rules, total


def build(rows: list[tuple], rules: list[str], total: float) -> Workbook:
    n = len(rows)
    orders = len({r[1] for r in rows})
    wb = Workbook()
    ws = wb.active
    ws.title = "Реестр Фурик"

    # ---- Титул-баннер ----
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOLS)
    put(
        ws,
        1,
        1,
        "РЕЕСТР ЛИЧНЫХ ПОКУПОК · СБЕРМАРКЕТ → ФУРКАТЖОН",
        fill=fill_header,
        font=font_title,
        border=border_accent,
        align=align_left,
    )
    ws.row_dimensions[1].height = 32

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NCOLS)
    put(
        ws,
        2,
        1,
        f"К перевыставлению  ·  {n} позиций  ·  {orders} заказов  ·  {date.today().strftime('%d.%m.%Y')}",
        fill=fill_header,
        font=font_sub,
        border=border_accent,
        align=align_left,
    )
    ws.row_dimensions[2].height = 20
    for c in range(1, NCOLS + 1):
        ws.cell(1, c).border = border_accent
        ws.cell(1, c).fill = fill_header
        ws.cell(2, c).border = border_accent
        ws.cell(2, c).fill = fill_header

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=NCOLS)
    put(
        ws,
        3,
        1,
        "Согласованные личные позиции из заказов доставки. Товары кофейни и сервисные сборы не включены.",
        fill=fill_zebra,
        font=Font(name=FONT_HEAD, size=11, bold=True, color=INK),
        align=align_left,
    )
    ws.row_dimensions[3].height = 22

    # ---- Мета Поле / Значение ----
    meta = [
        ("Кому", "Фуркатжон (перевыставление личных покупок)"),
        ("Источник", "Сбермаркет / Купер"),
        ("Позиций", str(n)),
        ("Заказов", str(orders)),
        ("Итого", f"{total:,.2f} ₽".replace(",", " ").replace(".", ",")),
    ]
    r0 = 5
    put(ws, r0, 1, "Поле", fill=fill_header, font=font_head, border=border_accent, align=align_center)
    put(ws, r0, 2, "Значение", fill=fill_header, font=font_head, border=border_accent, align=align_center)
    ws.merge_cells(start_row=r0, start_column=2, end_row=r0, end_column=4)
    for c in range(2, 5):
        ws.cell(r0, c).fill = fill_header
        ws.cell(r0, c).border = border_accent
        ws.cell(r0, c).font = font_head

    for i, (k, v) in enumerate(meta):
        r = r0 + 1 + i
        row_fill = fill_key if (i % 2) == 0 else fill_zebra
        put(ws, r, 1, k, fill=row_fill, font=font_body_bold, border=border_soft, align=align_left)
        put(ws, r, 2, v, fill=row_fill, font=font_body, border=border_soft, align=align_left)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        for c in range(2, 5):
            ws.cell(r, c).fill = row_fill
            ws.cell(r, c).border = border_soft

    # ---- Основная таблица ----
    hr = r0 + 1 + len(meta) + 2
    for c, h in enumerate(HEADERS, 1):
        put(
            ws,
            hr,
            c,
            h,
            fill=fill_header,
            font=font_head,
            border=border_accent,
            align=align_center,
        )
    ws.row_dimensions[hr].height = 28

    for i, row in enumerate(rows):
        r = hr + 1 + i
        row_fill = fill_key if (i % 2) == 0 else fill_zebra
        d, no, pos, qty, unit, sm, src = row
        vals = [d, no, pos, qty, unit, sm, src]
        for c, val in enumerate(vals, 1):
            if c == 1:
                align, font, nf = align_center, font_body, "DD.MM.YYYY"
            elif c == 2:
                align, font, nf = align_center, font_body_bold, None
            elif c == 3:
                align, font, nf = align_left, font_body, None
            elif c == 4:
                align, font, nf = align_center, font_body, "0"
            elif c in (5, 6):
                align, font, nf = align_right, font_body_bold if c == 6 else font_body, MONEY_FMT
            else:
                align, font, nf = align_left, font_body, None
            put(
                ws,
                r,
                c,
                val,
                fill=row_fill,
                font=font,
                border=border_soft,
                align=align,
                number_format=nf,
            )
        ws.row_dimensions[r].height = 22

    # ---- Итого ----
    tr = hr + 1 + n
    put(
        ws,
        tr,
        1,
        "",
        fill=fill_total,
        font=font_head,
        border=border_accent,
        align=align_left,
    )
    put(ws, tr, 2, "", fill=fill_total, font=font_head, border=border_accent)
    put(
        ws,
        tr,
        3,
        "Итого к перевыставлению",
        fill=fill_total,
        font=font_head,
        border=border_accent,
        align=align_left,
    )
    put(ws, tr, 4, "", fill=fill_total, font=font_head, border=border_accent)
    put(ws, tr, 5, "", fill=fill_total, font=font_head, border=border_accent)
    put(
        ws,
        tr,
        6,
        total,
        fill=fill_total,
        font=font_head,
        border=border_accent,
        align=align_right,
        number_format=MONEY_FMT,
    )
    put(
        ws,
        tr,
        7,
        f"{n} позиций",
        fill=fill_total,
        font=font_head,
        border=border_accent,
        align=align_left,
    )
    ws.row_dimensions[tr].height = 24

    widths = [12, 16, 48, 8, 12, 12, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = f"A{hr + 1}"
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(NCOLS)}{tr - 1}"
    ws.print_title_rows = f"{hr}:{hr}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---- Лист Правила ----
    leg = wb.create_sheet("Правила")
    leg.merge_cells("A1:B1")
    put(
        leg,
        1,
        1,
        "ПРАВИЛА ПЕРЕВЫСТАВЛЕНИЯ",
        fill=fill_header,
        font=font_title,
        border=border_accent,
        align=align_left,
    )
    leg.cell(1, 2).fill = fill_header
    leg.cell(1, 2).border = border_accent
    leg.row_dimensions[1].height = 30

    leg.merge_cells("A2:B2")
    put(
        leg,
        2,
        1,
        "Оформление как у договора ремонта (пастель: жёлтый / крем, текст чёрный)",
        fill=fill_header,
        font=font_sub,
        border=border_accent,
        align=align_left,
    )
    leg.cell(2, 2).fill = fill_header
    leg.cell(2, 2).border = border_accent

    put(leg, 4, 1, "№", fill=fill_header, font=font_head, border=border_accent, align=align_center)
    put(leg, 4, 2, "Правило", fill=fill_header, font=font_head, border=border_accent, align=align_center)

    body_rules = rules[1:] if rules and "что" in rules[0].lower() else rules
    if not body_rules:
        body_rules = [
            "Включать: зубные пасты, аджики, кетчупы, соусы, освежители воздуха, средства для стирки/гель, леденцы/конфеты Halls (+ иное по указанию).",
            "Не включать (кофейня): апельсины, грейпфрут, лайм, лимоны, салфетки, вода, мята, перчатки, тоники Rich, мёд, корица и пр. товар точки.",
            "Не включать сборку/доставку/сервисный сбор, пока не сказано иное.",
            "Дата = дата на карточке заказа в Купере (на скрине).",
            "Стоимость = цена за 1 шт; Сумма = стоимость × количество.",
        ]

    for i, text in enumerate(body_rules):
        r = 5 + i
        row_fill = fill_key if (i % 2) == 0 else fill_zebra
        put(leg, r, 1, i + 1, fill=row_fill, font=font_body_bold, border=border_soft, align=align_center)
        put(leg, r, 2, text, fill=row_fill, font=font_body, border=border_soft, align=align_left)
        leg.row_dimensions[r].height = 36

    note_r = 5 + len(body_rules) + 1
    leg.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=2)
    put(
        leg,
        note_r,
        1,
        f"Итого к перевыставлению: {total:,.2f} ₽ · {n} позиций.".replace(",", " ").replace(".", ","),
        fill=fill_total,
        font=font_head,
        border=border_accent,
        align=align_left,
    )
    leg.cell(note_r, 2).fill = fill_total
    leg.cell(note_r, 2).border = border_accent
    leg.row_dimensions[note_r].height = 28

    leg.column_dimensions["A"].width = 6
    leg.column_dimensions["B"].width = 100

    return wb


def main() -> None:
    rows, rules, total = load_data(SRC)
    assert len(rows) == 40, f"expected 40 rows, got {len(rows)}"
    assert abs(total - 9853.56) < 0.01, f"expected 9853.56, got {total}"
    wb = build(rows, rules, total)
    for out in OUTS:
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(out)
        print(out)
    print(f"OK: {len(rows)} rows, total={total}")


if __name__ == "__main__":
    main()
